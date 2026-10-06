from pathlib import Path

import pytest
from openpyxl import Workbook
from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QLabel,
    QPushButton,
    QTableWidget,
    QTabWidget,
)
from tests.application.data import (
    DIRECTORY_LINK,
    EMPTY_CARDS,
    FACT_BUYER_COLUMN,
    FACT_RULE_COLUMN,
    FIRST_SHOWN_ROW,
    HELD_RUN_STATUS,
    HELD_RUN_VAT,
    HOME_LINKS,
    KEPT_RUN,
    LOAD_LINK,
    LOAD_TITLES,
    MISMATCH_CALC_COLUMN,
    MISMATCH_FILE_COLUMN,
    MISMATCH_POINT_COLUMN,
    PLAN_POINT_COLUMN,
    SHOWN_MISMATCH_ROWS,
    STORED_PLAN_ROWS,
    TAB_TITLES,
    UNKNOWN_CONTRACT,
    UNKNOWN_FACT_GAPS,
    UNKNOWN_POINT,
)
from tests.ingest.data import (
    ACCEPTED_REMARKS,
    CALCULATED,
    EMPTY_RUNS,
    FIRST_MONTH,
    PARTY_ROWS,
    PLAN_HEADERS,
    PLAN_LINES,
    PLAN_ROWS,
    PLAN_UNWRITTEN,
    POINT,
    REJECTED_GROUP,
    STATED,
    TITLE,
    YEAR,
)
from tests.ingest.fact_data import ADDRESS, NEW_NAME, NEW_VOLUME
from tests.ingest.fact_data import TITLE as FACT_TITLE

from transport.application.catalog import Catalog
from transport.application.fact_load import NEW_CONSUMER
from transport.domain.group import Group
from transport.storage.repository import count_runs
from transport.ui.home import HomeWindow
from transport.ui.load import LoadWindow


@pytest.fixture
def qapp() -> QApplication:
    application = QApplication.instance()
    if application is None:
        application = QApplication([])
    return application


@pytest.mark.parametrize("chosen", (STATED, CALCULATED))
def test_plan_group_is_accepted_from_the_window(
    qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, chosen: Group
) -> None:
    catalog = Catalog.open(tmp_path / "plan.sqlite")
    try:
        window = LoadWindow(catalog)
        _choose_file(monkeypatch, _plan_book(tmp_path))
        _click(window, "loadPlan")
        table = _table(window, "mismatchTable")
        plan = _table(window, "planTable")
        assert plan.rowCount() == PLAN_LINES
        points = {_shown(plan, row, PLAN_POINT_COLUMN) for row in range(plan.rowCount())}
        assert POINT in points
        assert _label(window, "planMessage") == (
            f"План {YEAR} записан. Строк: {PLAN_LINES}. "
            f"Расхождений группы: {SHOWN_MISMATCH_ROWS}. Не записано: {PLAN_UNWRITTEN}."
        )
        assert table.rowCount() == SHOWN_MISMATCH_ROWS
        assert _shown(table, FIRST_SHOWN_ROW, MISMATCH_POINT_COLUMN) == POINT
        assert _shown(table, FIRST_SHOWN_ROW, MISMATCH_FILE_COLUMN) == STATED.value
        assert _shown(table, FIRST_SHOWN_ROW, MISMATCH_CALC_COLUMN) == CALCULATED.value
        assert len(catalog.consumers()) == PARTY_ROWS
        assert count_runs(catalog.connection()) == EMPTY_RUNS
        combo = window.findChild(QComboBox, "acceptGroup")
        assert combo is not None
        assert _has_group(combo, STATED)
        assert _has_group(combo, CALCULATED)
        assert not _has_group(combo, REJECTED_GROUP)
        _select_group(combo, chosen)
        _click(window, "acceptPlan")
        assert table.rowCount() == ACCEPTED_REMARKS
        _click(window, "loadPlan")
        assert table.rowCount() == ACCEPTED_REMARKS
        assert count_runs(catalog.connection()) == EMPTY_RUNS
    finally:
        catalog.close()


def test_reload_keeps_the_written_run(
    qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    catalog = Catalog.open(tmp_path / "kept.sqlite")
    try:
        window = LoadWindow(catalog)
        _choose_file(monkeypatch, _plan_book(tmp_path))
        _click(window, "loadPlan")
        assert count_runs(catalog.connection()) == EMPTY_RUNS
        catalog.connection().execute(
            "INSERT INTO run (year, month, status, with_vat) VALUES (?, ?, ?, ?)",
            (YEAR, FIRST_MONTH, HELD_RUN_STATUS, HELD_RUN_VAT),
        )
        catalog.connection().commit()
        _click(window, "loadPlan")
        assert count_runs(catalog.connection()) == KEPT_RUN
    finally:
        catalog.close()


def test_fact_discrepancies_do_not_create_cards(
    qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    catalog = Catalog.open(tmp_path / "fact.sqlite")
    try:
        window = LoadWindow(catalog)
        _choose_file(monkeypatch, _fact_book(tmp_path))
        _click(window, "loadFact")
        table = _table(window, "discrepancyTable")
        assert table.rowCount() == UNKNOWN_FACT_GAPS
        rules = {_shown(table, row, FACT_RULE_COLUMN) for row in range(table.rowCount())}
        names = {_shown(table, row, FACT_BUYER_COLUMN) for row in range(table.rowCount())}
        assert NEW_CONSUMER in rules
        assert NEW_NAME in names
        assert len(catalog.consumers()) == EMPTY_CARDS
        assert len(catalog.points()) == EMPTY_CARDS
        assert count_runs(catalog.connection()) == EMPTY_RUNS
    finally:
        catalog.close()


def test_home_opens_directories_and_load(qapp: QApplication, tmp_path: Path) -> None:
    catalog = Catalog.open(tmp_path / "home.sqlite")
    try:
        home = HomeWindow(catalog)
        assert _button_titles(home) == HOME_LINKS
        _click_text(home, DIRECTORY_LINK)
        directories = home.opened(DIRECTORY_LINK)
        assert directories is not None
        assert _tab_titles(directories, "directories") == TAB_TITLES
        _click_text(home, LOAD_LINK)
        load = home.opened(LOAD_LINK)
        assert load is not None
        assert _tab_titles(load, "loadTabs") == LOAD_TITLES
        assert _label(load, "planMessage") == f"В базе строк плана: {STORED_PLAN_ROWS}."
        _click_text(home, LOAD_LINK)
        assert home.opened(LOAD_LINK) is load
    finally:
        catalog.close()


def _choose_file(monkeypatch: pytest.MonkeyPatch, path: Path) -> None:
    monkeypatch.setattr(
        "transport.ui.load.QFileDialog.getOpenFileName",
        lambda *_args, **_kwargs: (str(path), ""),
    )


def _plan_book(directory: Path) -> Path:
    path = directory / "plan.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "план"
    sheet["A1"] = TITLE
    for column, value in enumerate(PLAN_HEADERS, start=1):
        sheet.cell(4, column, value)
    for offset, row in enumerate(PLAN_ROWS, start=5):
        for column, value in enumerate(row, start=1):
            sheet.cell(offset, column, value)
    book.save(path)
    book.close()
    return path


def _fact_book(directory: Path) -> Path:
    path = directory / "fact.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["B1"] = FACT_TITLE
    for column, value in enumerate(_FACT_HEADERS, start=1):
        sheet.cell(5, column, value)
    for column, value in enumerate(_FACT_SUB, start=1):
        sheet.cell(6, column, value)
    rows = (
        _buyer(NEW_NAME, UNKNOWN_CONTRACT),
        _city(UNKNOWN_CONTRACT, ADDRESS, UNKNOWN_POINT, NEW_VOLUME),
    )
    for offset, row in enumerate(rows, start=7):
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(offset, column, value)
    book.save(path)
    book.close()
    return path


def _buyer(name: str, contract: str) -> tuple[object, ...]:
    row: list[object] = [None] * len(_FACT_HEADERS)
    row[1] = name
    row[2] = contract
    return tuple(row)


def _city(contract: str, address: str, point: str, volume: object) -> tuple[object, ...]:
    row: list[object] = [None] * len(_FACT_HEADERS)
    row[2] = contract
    row[3] = volume
    row[4] = point
    row[5] = address
    row[6] = volume
    row[8] = volume
    return tuple(row)


def _table(window, name: str) -> QTableWidget:
    table = window.findChild(QTableWidget, name)
    assert table is not None
    return table


def _button_titles(window) -> tuple[str, ...]:
    layout = window.centralWidget().layout()
    assert layout is not None
    titles: list[str] = []
    for index in range(layout.count()):
        item = layout.itemAt(index)
        assert item is not None
        widget = item.widget()
        if isinstance(widget, QPushButton):
            titles.append(widget.text())
    return tuple(titles)


def _tab_titles(window, name: str) -> tuple[str, ...]:
    tabs = window.findChild(QTabWidget, name)
    assert tabs is not None
    return tuple(tabs.tabText(index) for index in range(tabs.count()))


def _click_text(window, title: str) -> None:
    for button in window.findChildren(QPushButton):
        if button.text() == title:
            button.click()
            return
    raise AssertionError(title)


def _label(window, name: str) -> str:
    label = window.findChild(QLabel, name)
    assert label is not None
    return label.text()


def _click(window, name: str) -> None:
    button = window.findChild(QPushButton, name)
    assert button is not None
    button.click()


def _shown(table: QTableWidget, row: int, column: int) -> str:
    item = table.item(row, column)
    assert item is not None
    return item.text()


def _has_group(combo: QComboBox, group: Group) -> bool:
    for index in range(combo.count()):
        if combo.itemData(index) == group.value:
            return True
    return False


def _select_group(combo: QComboBox, group: Group) -> None:
    for index in range(combo.count()):
        if combo.itemData(index) == group.value:
            combo.setCurrentIndex(index)
            return
    raise AssertionError(group.value)


_FACT_HEADERS = (
    "№ п\\п",
    "Покупатель",
    "Договор поставки газа",
    None,
    "Код ТП",
    "Адрес ТП",
    "Всего объем оттранспортированного газа, тыс. м3",
    None,
    "в т.ч. объем газа в пределах установленного по договору, тыс. м3",
    None,
    "в т.ч. объем газа сверхустановленного по договору (с 16 сентября по 14 апреля), тыс. м3",
    None,
    "в т.ч. объем газа сверхустановленного по договору (с 15 апреля по 15 сентября), тыс. м3",
    None,
    "Категория потребления",
    "Калорийность",
    "Тип газа",
)
_FACT_SUB = (None,) * 6 + ("в т.ч. СПб", "в т.ч. ЛО") * 4
