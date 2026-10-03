from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook
from tests.ingest.data import (
    ADDRESS,
    CALCULATED,
    CITY_CODE,
    CONSUMER_KIND,
    CONTRACT_A,
    CONTRACT_B,
    CONTRACT_ROWS,
    EMPTY_RUNS,
    FIRST_CODE,
    FIRST_MONTH,
    INN,
    MISMATCHES,
    MONTH,
    MONTH_COUNT,
    OBLAST_CODE,
    OBLAST_CONTRACT,
    OBLAST_GROUP,
    OBLAST_MONTH,
    OBLAST_VOLUME,
    OTHER_INN,
    OTHER_KIND,
    OTHER_POINT,
    PARTY_ROWS,
    PLAN_HEADERS,
    PLAN_LINES,
    PLAN_ROWS,
    POINT,
    POINT_ROWS,
    RAW_VOLUME,
    RAW_VOLUME_TEXT,
    SECOND_CODE,
    STATED,
    TITLE,
    VOLUME,
    YEAR,
)

from transport.application.plan_load import GROUP_MISMATCH, load_annual_plan
from transport.application.sample import CITY_CODE as SAMPLE_CITY
from transport.application.sample import OBLAST_CODE as SAMPLE_OBLAST
from transport.ingest.annual import PlanSheetError, read_annual_plan, volume_of
from transport.storage.database import open_database


def test_volume_rounds_binary_excel_number() -> None:
    assert volume_of(RAW_VOLUME) == RAW_VOLUME_TEXT


def test_directory_codes_are_the_plan_regions() -> None:
    assert SAMPLE_CITY == CITY_CODE
    assert SAMPLE_OBLAST == OBLAST_CODE


def test_total_row_is_not_a_plan_line(tmp_path: Path) -> None:
    book = read_annual_plan(_workbook(tmp_path))
    assert book.year == YEAR
    assert len(book.lines) == PLAN_LINES
    assert book.lines[0].volume == VOLUME
    assert book.lines[0].months[0] == MONTH
    assert book.lines[0].stated == STATED
    assert len(book.issues) == 1


def test_load_splits_contracts_and_sums_the_group(tmp_path: Path) -> None:
    path = _workbook(tmp_path)
    connection = open_database(tmp_path / "base.sqlite")
    try:
        loaded = load_annual_plan(connection, path)
        again = load_annual_plan(connection, path)
        assert loaded.year == YEAR
        assert loaded.lines == PLAN_LINES
        assert again.lines == PLAN_LINES
        assert loaded.mismatches == ((POINT, STATED, CALCULATED), (POINT, STATED, CALCULATED))
        assert _count(connection, "consumer") == PARTY_ROWS
        assert _count(connection, "contract") == CONTRACT_ROWS
        assert _count(connection, "point") == POINT_ROWS
        assert _count(connection, "annual_plan") == PLAN_LINES
        assert _count(connection, "monthly_plan") == PLAN_LINES * MONTH_COUNT
        assert _count(connection, "run") == EMPTY_RUNS
        assert _codes(connection) == (FIRST_CODE, SECOND_CODE)
        assert _point_contract(connection, POINT) == _contract_id(connection, CONTRACT_A)
        assert _plan_volume(connection, POINT, CONTRACT_B) == VOLUME
        assert _plan_volume(connection, OTHER_POINT, OBLAST_CONTRACT) == OBLAST_VOLUME
        assert _month_volume(connection, POINT, CONTRACT_A) == MONTH
        assert _month_volume(connection, OTHER_POINT, OBLAST_CONTRACT) == OBLAST_MONTH
        assert _group(connection, POINT) == CALCULATED.value
        assert _group(connection, OTHER_POINT) == OBLAST_GROUP.value
        assert _kind(connection, INN) == CONSUMER_KIND.value
        assert _kind(connection, OTHER_INN) == OTHER_KIND.value
        assert _region(connection, POINT, CONTRACT_A) == CITY_CODE
        assert _region(connection, OTHER_POINT, OBLAST_CONTRACT) == OBLAST_CODE
        assert _address(connection, POINT) == ADDRESS
        remarks = connection.execute(
            "SELECT COUNT(*) FROM remark WHERE rule_code = ?",
            (GROUP_MISMATCH,),
        ).fetchone()
        assert remarks == (MISMATCHES,)
    finally:
        connection.close()


def test_missing_sheet_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    book = Workbook()
    book.save(path)
    book.close()
    with pytest.raises(PlanSheetError):
        read_annual_plan(path)


def _workbook(directory: Path) -> Path:
    path = directory / "plan.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "план"
    sheet["A1"] = TITLE
    for column, value in enumerate(PLAN_HEADERS, start=1):
        sheet.cell(4, column, value)
    for offset, row in enumerate(PLAN_ROWS, start=5):
        for column, value in enumerate(row, start=1):
            sheet.cell(offset, column, value)
    book.save(path)
    book.close()
    return path


def _count(connection, table: str) -> int:
    row = connection.execute(_COUNT[table]).fetchone()
    assert row is not None
    return int(row[0])


def _codes(connection) -> tuple[str, ...]:
    rows = connection.execute("SELECT code FROM consumer ORDER BY code").fetchall()
    return tuple(str(row[0]) for row in rows)


def _contract_id(connection, number: str) -> int:
    row = connection.execute("SELECT id FROM contract WHERE number = ?", (number,)).fetchone()
    assert row is not None
    return int(row[0])


def _point_contract(connection, code: str) -> int:
    row = connection.execute("SELECT contract_id FROM point WHERE code = ?", (code,)).fetchone()
    assert row is not None
    return int(row[0])


def _plan_volume(connection, point: str, contract: str) -> Decimal:
    row = connection.execute(
        """
        SELECT annual_plan.volume
        FROM annual_plan
        JOIN point ON point.id = annual_plan.point_id
        JOIN contract ON contract.id = annual_plan.contract_id
        WHERE point.code = ? AND contract.number = ?
        """,
        (point, contract),
    ).fetchone()
    assert row is not None
    return Decimal(str(row[0]))


def _month_volume(connection, point: str, contract: str) -> Decimal:
    row = connection.execute(
        """
        SELECT monthly_plan.volume
        FROM monthly_plan
        JOIN point ON point.id = monthly_plan.point_id
        JOIN contract ON contract.id = monthly_plan.contract_id
        WHERE point.code = ? AND contract.number = ? AND monthly_plan.month = ?
        """,
        (point, contract, FIRST_MONTH),
    ).fetchone()
    assert row is not None
    return Decimal(str(row[0]))


def _group(connection, code: str) -> str:
    row = connection.execute(
        """
        SELECT group_code FROM point_group
        JOIN point ON point.id = point_group.point_id
        WHERE point.code = ?
        """,
        (code,),
    ).fetchone()
    assert row is not None
    return str(row[0])


def _kind(connection, inn: str) -> str:
    row = connection.execute("SELECT kind FROM consumer WHERE inn = ?", (inn,)).fetchone()
    assert row is not None
    return str(row[0])


def _region(connection, point: str, contract: str) -> str:
    row = connection.execute(
        """
        SELECT region.code
        FROM annual_plan
        JOIN region ON region.id = annual_plan.region_id
        JOIN point ON point.id = annual_plan.point_id
        JOIN contract ON contract.id = annual_plan.contract_id
        WHERE point.code = ? AND contract.number = ?
        """,
        (point, contract),
    ).fetchone()
    assert row is not None
    return str(row[0])


def _address(connection, code: str) -> str:
    row = connection.execute("SELECT address FROM point WHERE code = ?", (code,)).fetchone()
    assert row is not None
    return str(row[0])


_COUNT = {
    "consumer": "SELECT COUNT(*) FROM consumer",
    "contract": "SELECT COUNT(*) FROM contract",
    "point": "SELECT COUNT(*) FROM point",
    "annual_plan": "SELECT COUNT(*) FROM annual_plan",
    "monthly_plan": "SELECT COUNT(*) FROM monthly_plan",
    "run": "SELECT COUNT(*) FROM run",
}
