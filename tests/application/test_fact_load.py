from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from tests.ingest.fact_data import (
    ADDRESS,
    APRIL_TRANSITIONS,
    CHEAPER,
    CONTRACT_LO,
    CONTRACT_NAME,
    CONTRACT_OK,
    CONTRACT_PLAN,
    CROSS_CONTRACT,
    CROSS_CONTRACT_B,
    CROSS_FROM,
    CROSS_MONTH,
    CROSS_POINT,
    CROSS_PRIOR,
    CROSS_TO,
    CROSS_VOLUME,
    DEAR_CONTRACT,
    DEAR_FROM,
    DEAR_POINT,
    DEAR_TO,
    DEAR_VOLUME,
    DEARER,
    DECEMBER,
    DECEMBER_TITLE,
    DECEMBER_TRANSITIONS,
    DIRECTORY_CONSUMERS,
    DISCREPANCY_COUNTS,
    DISCREPANCY_ROWS,
    EMPTY_RUNS,
    FACT_ROWS,
    FACT_VOLUME,
    FILE_ADDRESS,
    FIRST_CODE,
    HELD_LINES,
    HOLD_CONTRACT,
    HOLD_OVER,
    HOLD_POINT,
    HOLD_PRIOR,
    JANUARY,
    KNOWN_NEW_CONTRACT,
    KNOWN_NEW_POINT,
    LO_110,
    LO_150,
    LO_ADDRESS,
    LO_NAME,
    LO_VOLUME,
    MISMATCH_NAME,
    MONTH,
    NAME,
    NEW_CONSUMER,
    NEW_CONTRACT,
    NEW_NAME,
    NEW_POINT,
    NEW_VOLUME,
    OBLAST_CODE,
    OBLAST_NAME,
    OVER_110,
    OVER_150,
    PARSED_LINES,
    PLAN_NAME,
    PLAN_VOLUME,
    POINT_ADDRESS,
    POINT_LO,
    POINT_NAME,
    POINT_OK,
    POINT_PLAN,
    POINT_ZERO,
    REGION_CODE,
    REGION_NAME,
    REPLACEMENT_VOLUME,
    SECOND_CODE,
    SIGNED_ON,
    SPACED_NAME,
    STORED_ADDRESS,
    THIRD_CODE,
    TITLE,
    TRANSITION_NAME,
    YEAR,
    ZERO_VOLUME,
)

from transport.application.fact_load import load_month_fact
from transport.domain.group import Group
from transport.domain.month import ConsumerKind
from transport.ingest.fact import read_month_fact
from transport.storage.database import open_database
from transport.storage.repository import (
    save_consumer,
    save_contract,
    save_monthly_fact,
    save_monthly_plan,
    save_point,
    save_point_group,
    save_region,
)


def test_month_fact_records_gaps_and_skips_new_cards(tmp_path: Path) -> None:
    connection = open_database(tmp_path / "base.sqlite")
    try:
        _directory(connection)
        path = _workbook(tmp_path, FACT_VOLUME)
        parsed = read_month_fact(path)
        loaded = load_month_fact(connection, path)
        assert parsed.year == YEAR
        assert parsed.month == MONTH
        assert len(parsed.lines) == PARSED_LINES
        assert loaded.lines == FACT_ROWS
        assert loaded.discrepancies == DISCREPANCY_ROWS
        assert _count(connection, "SELECT COUNT(*) FROM consumer") == DIRECTORY_CONSUMERS
        assert _count(connection, "SELECT COUNT(*) FROM monthly_fact") == FACT_ROWS
        assert _count(connection, "SELECT COUNT(*) FROM run") == EMPTY_RUNS
        assert _fact(connection, CONTRACT_OK, POINT_OK) == (FACT_VOLUME, OVER_110, OVER_150)
        assert _fact(connection, CONTRACT_LO, POINT_LO) == (LO_VOLUME, LO_110, LO_150)
        assert _fact(connection, NEW_CONTRACT, NEW_POINT) is None
        assert _address(connection, POINT_ADDRESS) == STORED_ADDRESS
        assert _name(connection, CONTRACT_NAME) == NAME
        for code, count in DISCREPANCY_COUNTS:
            assert _rules(connection, code) == count
        connection.execute(
            """
            INSERT INTO fact_discrepancy (year, month, rule_code, message)
            VALUES (?, ?, ?, ?)
            """,
            (YEAR, MONTH, NEW_CONSUMER, "чужая"),
        )
        again = load_month_fact(connection, _workbook(tmp_path, REPLACEMENT_VOLUME))
        assert again.lines == FACT_ROWS
        assert again.discrepancies == DISCREPANCY_ROWS
        assert _fact(connection, CONTRACT_OK, POINT_OK) == (REPLACEMENT_VOLUME, OVER_110, OVER_150)
        assert _rules(connection, NEW_CONSUMER) == dict(DISCREPANCY_COUNTS)[NEW_CONSUMER]
        assert _count(connection, "SELECT COUNT(*) FROM fact_discrepancy") == DISCREPANCY_ROWS
    finally:
        connection.close()


def test_empty_point_contract_comes_from_the_buyer(tmp_path: Path) -> None:
    book = read_month_fact(_held(tmp_path))
    assert len(book.lines) == HELD_LINES
    assert book.lines[0].contract == CONTRACT_OK
    assert book.lines[0].point == POINT_OK
    assert book.lines[0].name == NAME
    assert book.lines[0].volume == FACT_VOLUME


def test_fact_load_lists_points_that_change_group(tmp_path: Path) -> None:
    connection = open_database(tmp_path / "transition.sqlite")
    try:
        cross, hold, dear = _transition_directory(connection)
        load_month_fact(connection, _transition_book(tmp_path, TITLE, _april_rows()))
        assert _transition(connection, MONTH, CROSS_POINT) == (
            CROSS_FROM,
            CROSS_TO,
            CROSS_VOLUME,
            CHEAPER,
        )
        assert _transition(connection, MONTH, HOLD_POINT) is None
        assert _transition(connection, MONTH, DEAR_POINT) is None
        assert _count(connection, "SELECT COUNT(*) FROM group_transition") == APRIL_TRANSITIONS
        assert _group(connection, CROSS_POINT) == CROSS_FROM.value
        load_month_fact(
            connection,
            _transition_book(tmp_path, DECEMBER_TITLE, _december_rows()),
        )
        assert _transition(connection, DECEMBER, CROSS_POINT) == (
            CROSS_FROM,
            CROSS_TO,
            CROSS_VOLUME,
            CHEAPER,
        )
        assert _transition(connection, DECEMBER, DEAR_POINT) == (
            DEAR_FROM,
            DEAR_TO,
            DEAR_VOLUME,
            DEARER,
        )
        assert _count(connection, "SELECT COUNT(*) FROM group_transition") == (
            APRIL_TRANSITIONS + DECEMBER_TRANSITIONS
        )
        assert _group(connection, dear) == DEAR_FROM.value
        assert _group(connection, cross) == CROSS_FROM.value
        assert _group(connection, hold) == CROSS_FROM.value
        assert _count(connection, "SELECT COUNT(*) FROM run") == EMPTY_RUNS
    finally:
        connection.close()


def _transition_directory(connection) -> tuple[int, int, int]:
    region_id = save_region(connection, code=REGION_CODE, name=REGION_NAME)
    consumer_id = _consumer(connection, region_id, FIRST_CODE, TRANSITION_NAME)
    first = _contract(connection, consumer_id, CROSS_CONTRACT)
    second = _contract(connection, consumer_id, CROSS_CONTRACT_B)
    hold = _contract(connection, consumer_id, HOLD_CONTRACT)
    dear = _contract(connection, consumer_id, DEAR_CONTRACT)
    cross_id = _point(connection, first, CROSS_POINT, ADDRESS)
    hold_id = _point(connection, hold, HOLD_POINT, ADDRESS)
    dear_id = _point(connection, dear, DEAR_POINT, ADDRESS)
    _group_row(connection, cross_id, CROSS_FROM)
    _group_row(connection, hold_id, CROSS_FROM)
    _group_row(connection, dear_id, DEAR_FROM)
    _prior_fact(connection, first, cross_id, JANUARY, CROSS_PRIOR, ZERO_VOLUME, ZERO_VOLUME)
    _prior_fact(connection, second, cross_id, JANUARY, CROSS_PRIOR, ZERO_VOLUME, ZERO_VOLUME)
    _prior_fact(connection, hold, hold_id, JANUARY, HOLD_PRIOR, ZERO_VOLUME, ZERO_VOLUME)
    return cross_id, hold_id, dear_id


def _group_row(connection, point_id: int, group: Group) -> None:
    save_point_group(connection, point_id=point_id, effective_from=SIGNED_ON, group=group)


def _prior_fact(
    connection,
    contract_id: int,
    point_id: int,
    month: int,
    volume: Decimal,
    winter: Decimal,
    summer: Decimal,
) -> None:
    save_monthly_fact(
        connection,
        contract_id=contract_id,
        point_id=point_id,
        year=YEAR,
        month=month,
        volume=volume,
        overlimit_110=summer,
        overlimit_150=winter,
        kind=ConsumerKind.INDUSTRIAL,
    )


def _april_rows() -> tuple[tuple[object, ...], ...]:
    return (
        _buyer(1, TRANSITION_NAME, CROSS_CONTRACT),
        _city(CROSS_CONTRACT, ADDRESS, CROSS_POINT, CROSS_MONTH, ZERO_VOLUME, ZERO_VOLUME),
        _buyer(2, TRANSITION_NAME, HOLD_CONTRACT),
        _city(HOLD_CONTRACT, ADDRESS, HOLD_POINT, CROSS_MONTH, HOLD_OVER, ZERO_VOLUME),
    )


def _december_rows() -> tuple[tuple[object, ...], ...]:
    return (
        _buyer(1, TRANSITION_NAME, DEAR_CONTRACT),
        _city(DEAR_CONTRACT, ADDRESS, DEAR_POINT, DEAR_VOLUME, ZERO_VOLUME, ZERO_VOLUME),
    )


def _transition_book(directory: Path, title: str, rows: tuple[tuple[object, ...], ...]) -> Path:
    path = directory / "transition.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["B1"] = title
    for column, value in enumerate(_HEADERS, start=1):
        sheet.cell(5, column, value)
    for column, value in enumerate(_SUB, start=1):
        sheet.cell(6, column, value)
    for offset, row in enumerate(rows, start=7):
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(offset, column, value)
    book.save(path)
    book.close()
    return path


def _transition(connection, month: int, point: str) -> tuple[Group, Group, Decimal, str] | None:
    row = connection.execute(
        """
        SELECT recorded_group, calculated_group, volume, direction
        FROM group_transition
        JOIN point ON point.id = group_transition.point_id
        WHERE group_transition.month = ? AND point.code = ?
        """,
        (month, point),
    ).fetchone()
    if row is None:
        return None
    return Group(str(row[0])), Group(str(row[1])), Decimal(str(row[2])), str(row[3])


def _group(connection, point: int | str) -> str:
    if isinstance(point, int):
        row = connection.execute(
            "SELECT group_code FROM point_group WHERE point_id = ?",
            (point,),
        ).fetchone()
    else:
        row = connection.execute(
            """
            SELECT group_code FROM point_group
            JOIN point ON point.id = point_group.point_id
            WHERE point.code = ?
            """,
            (point,),
        ).fetchone()
    assert row is not None
    return str(row[0])


def _directory(connection) -> None:
    region_id = save_region(connection, code=REGION_CODE, name=REGION_NAME)
    save_region(connection, code=OBLAST_CODE, name=OBLAST_NAME)
    plant = _consumer(connection, region_id, FIRST_CODE, NAME)
    village = _consumer(connection, region_id, SECOND_CODE, PLAN_NAME)
    oblast = _consumer(connection, region_id, THIRD_CODE, LO_NAME)
    ok = _contract(connection, plant, CONTRACT_OK)
    named = _contract(connection, plant, CONTRACT_NAME)
    planned = _contract(connection, village, CONTRACT_PLAN)
    lo = _contract(connection, oblast, CONTRACT_LO)
    _point(connection, ok, POINT_OK, ADDRESS)
    _point(connection, ok, POINT_ADDRESS, STORED_ADDRESS)
    _point(connection, named, POINT_NAME, ADDRESS)
    _point(connection, planned, POINT_PLAN, ADDRESS)
    _point(connection, planned, POINT_ZERO, ADDRESS)
    _point(connection, lo, POINT_LO, LO_ADDRESS)
    save_monthly_plan(
        connection,
        contract_id=planned,
        point_id=_point_id(connection, POINT_PLAN),
        year=YEAR,
        month=MONTH,
        volume=PLAN_VOLUME,
    )
    save_monthly_plan(
        connection,
        contract_id=planned,
        point_id=_point_id(connection, POINT_ZERO),
        year=YEAR,
        month=MONTH,
        volume=ZERO_VOLUME,
    )


def _consumer(connection, region_id: int, code: str, name: str) -> int:
    return save_consumer(
        connection,
        code=code,
        name=name,
        region_id=region_id,
        kind=ConsumerKind.INDUSTRIAL,
        inn=None,
        on=SIGNED_ON,
    )


def _contract(connection, consumer_id: int, number: str) -> int:
    return save_contract(
        connection,
        consumer_id=consumer_id,
        number=number,
        signed_on=SIGNED_ON,
        on=SIGNED_ON,
    )


def _point(connection, contract_id: int, code: str, address: str) -> int:
    return save_point(
        connection,
        contract_id=contract_id,
        code=code,
        address=address,
        on=SIGNED_ON,
    )


def _held(directory: Path) -> Path:
    path = directory / "held.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["B1"] = TITLE
    for column, value in enumerate(_HEADERS, start=1):
        sheet.cell(5, column, value)
    for column, value in enumerate(_SUB, start=1):
        sheet.cell(6, column, value)
    point = list(_city(CONTRACT_OK, ADDRESS, POINT_OK, FACT_VOLUME, OVER_150, OVER_110))
    point[2] = None
    rows = (_buyer(1, NAME, CONTRACT_OK), tuple(point))
    for offset, row in enumerate(rows, start=7):
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(offset, column, value)
    book.save(path)
    book.close()
    return path


def _workbook(directory: Path, volume: Decimal) -> Path:
    path = directory / "fact.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Лист1"
    sheet["B1"] = TITLE
    for column, value in enumerate(_HEADERS, start=1):
        sheet.cell(5, column, value)
    for column, value in enumerate(_SUB, start=1):
        sheet.cell(6, column, value)
    rows = (
        _buyer(1, SPACED_NAME, CONTRACT_OK),
        _city(CONTRACT_OK, ADDRESS, POINT_OK, volume, OVER_150, OVER_110),
        _city(CONTRACT_OK, FILE_ADDRESS, POINT_ADDRESS, volume, OVER_150, OVER_110),
        _buyer(2, MISMATCH_NAME, CONTRACT_NAME),
        _city(CONTRACT_NAME, ADDRESS, POINT_NAME, volume, OVER_150, OVER_110),
        _buyer(3, LO_NAME, CONTRACT_LO),
        _oblast(CONTRACT_LO, LO_ADDRESS, POINT_LO, LO_VOLUME),
        _buyer(4, NEW_NAME, NEW_CONTRACT),
        _city(NEW_CONTRACT, ADDRESS, NEW_POINT, NEW_VOLUME, ZERO_VOLUME, ZERO_VOLUME),
        _buyer(5, NAME, KNOWN_NEW_CONTRACT),
        _city(KNOWN_NEW_CONTRACT, ADDRESS, KNOWN_NEW_POINT, NEW_VOLUME, ZERO_VOLUME, ZERO_VOLUME),
    )
    for offset, row in enumerate(rows, start=7):
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(offset, column, value)
    book.save(path)
    book.close()
    return path


def _buyer(number: int, name: str, contract: str) -> tuple[object, ...]:
    row: list[object] = [None] * 17
    row[0] = number
    row[1] = name
    row[2] = contract
    return tuple(row)


def _city(
    contract: str,
    address: str,
    point: str,
    volume: Decimal,
    winter: Decimal,
    summer: Decimal,
) -> tuple[object, ...]:
    row: list[object] = [None] * 17
    row[1] = address
    row[2] = contract
    row[3] = volume
    row[4] = point
    row[5] = address
    row[6] = volume
    row[8] = volume - winter - summer
    row[10] = winter
    row[12] = summer
    return tuple(row)


def _oblast(contract: str, address: str, point: str, volume: Decimal) -> tuple[object, ...]:
    row: list[object] = [None] * 17
    row[1] = address
    row[2] = contract
    row[3] = volume
    row[4] = point
    row[5] = address
    row[7] = volume
    row[9] = volume
    return tuple(row)


def _fact(connection, contract: str, point: str) -> tuple[Decimal, Decimal, Decimal] | None:
    row = connection.execute(
        """
        SELECT monthly_fact.volume, monthly_fact.overlimit_110, monthly_fact.overlimit_150
        FROM monthly_fact
        JOIN contract ON contract.id = monthly_fact.contract_id
        JOIN point ON point.id = monthly_fact.point_id
        WHERE contract.number = ? AND point.code = ?
        """,
        (contract, point),
    ).fetchone()
    if row is None:
        return None
    return Decimal(str(row[0])), Decimal(str(row[1])), Decimal(str(row[2]))


def _rules(connection, code: str) -> int:
    row = connection.execute(
        "SELECT COUNT(*) FROM fact_discrepancy WHERE rule_code = ?",
        (code,),
    ).fetchone()
    assert row is not None
    return int(row[0])


def _count(connection, query: str) -> int:
    row = connection.execute(query).fetchone()
    assert row is not None
    return int(row[0])


def _address(connection, code: str) -> str:
    row = connection.execute("SELECT address FROM point WHERE code = ?", (code,)).fetchone()
    assert row is not None
    return str(row[0])


def _name(connection, number: str) -> str:
    row = connection.execute(
        """
        SELECT consumer.name
        FROM contract
        JOIN consumer ON consumer.id = contract.consumer_id
        WHERE contract.number = ?
        """,
        (number,),
    ).fetchone()
    assert row is not None
    return str(row[0])


def _point_id(connection, code: str) -> int:
    row = connection.execute("SELECT id FROM point WHERE code = ?", (code,)).fetchone()
    assert row is not None
    return int(row[0])


_HEADERS = (
    "№ п\\п",
    "Покупатель",
    "Договор поставки газа",
    None,
    "Код ТП",
    "Адрес ТП",
    "Всего объем оттранспортированного газа, тыс. м3",
    None,
    "в т.ч. объем газа в пределах установленного по договору, тыс. м3",
    None,
    "в т.ч. объем газа сверхустановленного по договору (с 16 сентября по 14 апреля), тыс. м3",
    None,
    "в т.ч. объем газа сверхустановленного по договору (с 15 апреля по 15 сентября), тыс. м3",
    None,
    "Категория потребления",
    "Калорийность",
    "Тип газа",
)
_SUB = (None,) * 6 + ("в т.ч. СПб", "в т.ч. ЛО") * 4
