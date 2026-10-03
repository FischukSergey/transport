"""Лист факта одного месяца.

Строка покупателя без кода точки в факт не входит: это сумма его точек.
Категория, калорийность и тип газа не читаются.
Год и месяц берутся из заголовка листа.
"""

import re
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from transport.ingest.annual import VOLUME_PLACES, volume_of

# 16 сентября — 14 апреля, коэффициент 1,5. 15 апреля — 15 сентября, коэффициент 1,1.
_WINTER = "16 сентября"
_SUMMER = "15 апреля"
_TOTAL = "всего объем"
_ZERO = Decimal(0).quantize(VOLUME_PLACES)

_MONTHS = {
    "январь": 1,
    "февраль": 2,
    "март": 3,
    "апрель": 4,
    "май": 5,
    "июнь": 6,
    "июль": 7,
    "август": 8,
    "сентябрь": 9,
    "октябрь": 10,
    "ноябрь": 11,
    "декабрь": 12,
}
_YEAR = re.compile(r"(20\d{2})")


class FactSheetError(Exception):
    """В книге нет листа факта с годом и месяцем. Строки не читаются."""

    def __init__(self, path: Path) -> None:
        self.path = path
        super().__init__(str(path))


class FactLine:
    def __init__(
        self,
        file_row: int,
        name: str,
        contract: str,
        point: str,
        address: str,
        volume: Decimal,
        overlimit_110: Decimal,
        overlimit_150: Decimal,
    ) -> None:
        self.file_row = file_row
        self.name = name
        self.contract = contract
        self.point = point
        self.address = address
        self.volume = volume
        self.overlimit_110 = overlimit_110
        self.overlimit_150 = overlimit_150


class FactFile:
    def __init__(self, year: int, month: int, lines: tuple[FactLine, ...]) -> None:
        self.year = year
        self.month = month
        self.lines = lines


class _Columns:
    def __init__(
        self,
        buyer: int,
        contract: int,
        point: int,
        address: int,
        total: tuple[int, int],
        winter: tuple[int, int],
        summer: tuple[int, int],
    ) -> None:
        self.buyer = buyer
        self.contract = contract
        self.point = point
        self.address = address
        self.total = total
        self.winter = winter
        self.summer = summer


def read_month_fact(path: Path | str) -> FactFile:
    """Читает факт месяца. В базу не пишет."""
    source = Path(path)
    book = load_workbook(source, data_only=True)
    try:
        grid = _grid(book)
    finally:
        book.close()
    if grid is None:
        raise FactSheetError(source)
    return _parse(source, grid)


def _grid(book: object) -> list[tuple[object, ...]] | None:
    sheets = getattr(book, "worksheets", ())
    for sheet in sheets:
        grid = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        if _header_row(grid) is not None:
            return grid
    return None


def _parse(path: Path, grid: list[tuple[object, ...]]) -> FactFile:
    header_at = _header_row(grid)
    if header_at is None or header_at + 1 >= len(grid):
        raise FactSheetError(path)
    period = _period(grid[:header_at])
    if period is None:
        raise FactSheetError(path)
    year, month = period
    columns = _columns(grid[header_at], grid[header_at + 1])
    if columns is None:
        raise FactSheetError(path)
    lines: list[FactLine] = []
    buyer = ""
    for offset, row in enumerate(grid[header_at + 2 :], start=header_at + 3):
        point = _text(row, columns.point)
        contract = _text(row, columns.contract)
        if point == "":
            name = _text(row, columns.buyer)
            if contract != "" and name != "" and not name.casefold().startswith("итого"):
                buyer = name
            continue
        if point.casefold().startswith("итого"):
            continue
        lines.append(
            FactLine(
                offset,
                buyer,
                contract,
                point,
                _text(row, columns.address),
                _pair(row, columns.total),
                _pair(row, columns.summer),
                _pair(row, columns.winter),
            )
        )
    return FactFile(year, month, tuple(lines))


def _header_row(grid: list[tuple[object, ...]]) -> int | None:
    for index, row in enumerate(grid):
        if any(_key(cell) == "код тп" for cell in row):
            return index
    return None


def _period(rows: list[tuple[object, ...]]) -> tuple[int, int] | None:
    for row in rows:
        for cell in row:
            text = _key(cell)
            if text == "":
                continue
            found = _YEAR.search(text)
            month = next((number for name, number in _MONTHS.items() if name in text), None)
            if found is not None and month is not None:
                return int(found.group(1)), month
    return None


def _columns(header: tuple[object, ...], sub: tuple[object, ...]) -> _Columns | None:
    buyer = contract = point = address = None
    total = winter = summer = None
    for index, cell in enumerate(header):
        key = _key(cell)
        if key == "покупатель":
            buyer = index
        elif key == "договор поставки газа":
            contract = index
        elif key == "код тп":
            point = index
        elif key == "адрес тп":
            address = index
        elif _TOTAL in key:
            total = (index, index + 1)
        elif _WINTER in key:
            winter = (index, index + 1)
        elif _SUMMER in key:
            summer = (index, index + 1)
    if (
        buyer is None
        or contract is None
        or point is None
        or address is None
        or total is None
        or winter is None
        or summer is None
    ):
        return None
    if _index(sub, total[0]) != "в т.ч. спб" or _index(sub, total[1]) != "в т.ч. ло":
        return None
    return _Columns(buyer, contract, point, address, total, winter, summer)


def _index(row: tuple[object, ...], column: int) -> str:
    if column >= len(row):
        return ""
    return _key(row[column])


def _pair(row: tuple[object, ...], columns: tuple[int, int]) -> Decimal:
    left, right = columns
    return _volume(_cell(row, left)) + _volume(_cell(row, right))


def _volume(value: object) -> Decimal:
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return _ZERO
    return volume_of(value)


def _text(row: tuple[object, ...], column: int) -> str:
    value = _cell(row, column)
    if value is None:
        return ""
    return str(value).strip()


def _cell(row: tuple[object, ...], column: int) -> object:
    if column >= len(row):
        return None
    return row[column]


def _key(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("ё", "е").replace("Ё", "Е")
    return " ".join(text.casefold().split())
