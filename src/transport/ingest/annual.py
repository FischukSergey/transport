"""Годовой план.

Лист ищется по заголовку кода точки: имя листа может быть не «план».
Строка «Итого» в план не входит. Календарный год берётся из шапки листа.
Колонка объёма года — «Год» или «2026 год». Суточный минимум и максимум не читаются.
"""

import re
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from transport.domain.group import Group
from transport.domain.month import ConsumerKind

# Объём плана — тыс. м³, три знака.
VOLUME_PLACES = Decimal("0.001")

_SHEET = "план"
_YEAR = re.compile(r"(20\d{2})\s+год")
_MONTHS = (
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
)
_REGION_CODES = {
    "санкт-петербург": "78",
    "ленинградская область": "47",
}
_TOTAL = "итого"


class PlanSheetError(Exception):
    """В книге нет листа плана или года в заголовке. Строки не разбираются."""

    def __init__(self, path: Path) -> None:
        self.path = path
        super().__init__(path.name)


class PlanLine:
    def __init__(
        self,
        file_row: int,
        name: str,
        inn: str,
        contract: str,
        point: str,
        region_code: str,
        address: str,
        kind: ConsumerKind,
        stated: Group,
        volume: Decimal,
        months: tuple[Decimal, ...],
    ) -> None:
        self.file_row = file_row
        self.name = name
        self.inn = inn
        self.contract = contract
        self.point = point
        self.region_code = region_code
        self.address = address
        self.kind = kind
        self.stated = stated
        self.volume = volume
        self.months = months


class PlanIssue:
    def __init__(self, file_row: int, text: str) -> None:
        self.file_row = file_row
        self.text = text


class PlanFile:
    def __init__(
        self,
        year: int,
        lines: tuple[PlanLine, ...],
        issues: tuple[PlanIssue, ...],
    ) -> None:
        self.year = year
        self.lines = lines
        self.issues = issues


def read_annual_plan(path: Path | str) -> PlanFile:
    """Читает лист годового плана. В базу не пишет.

    Сначала берётся лист «план», если на нём есть код точки. Иначе — первый
    лист с таким заголовком.
    """
    source = Path(path)
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        grid = _plan_grid(book)
    finally:
        book.close()
    if grid is None:
        raise PlanSheetError(source)
    return _parse(source, grid)


def _plan_grid(book: Workbook) -> list[tuple[object, ...]] | None:
    names = list(book.sheetnames)
    if _SHEET in names:
        names.remove(_SHEET)
        names.insert(0, _SHEET)
    for name in names:
        grid = [tuple(row) for row in book[name].iter_rows(values_only=True)]
        if _header_row(grid) is not None:
            return grid
    return None


def volume_of(value: object) -> Decimal:
    """Переводит ячейку Excel в тыс. м³ с тремя знаками.

    Двоичное число берётся кратчайшей записью и затем округляется.
    """
    if isinstance(value, bool) or value is None:
        raise ValueError(value)
    if isinstance(value, Decimal):
        number = value
    elif isinstance(value, int):
        number = Decimal(value)
    elif isinstance(value, float):
        number = Decimal(str(value))
    else:
        body = str(value).strip().replace(" ", "").replace("\u00a0", "").replace(",", ".")
        number = Decimal(body)
    return number.quantize(VOLUME_PLACES, rounding=ROUND_HALF_UP)


def _parse(path: Path, grid: list[tuple[object, ...]]) -> PlanFile:
    header_at = _header_row(grid)
    if header_at is None:
        raise PlanSheetError(path)
    year = _year_of(grid[:header_at], grid[header_at])
    if year is None:
        raise PlanSheetError(path)
    columns = _columns(grid[header_at])
    lines: list[PlanLine] = []
    issues: list[PlanIssue] = []
    for offset, row in enumerate(grid[header_at + 1 :], start=header_at + 2):
        if _is_total(row, columns):
            continue
        if _blank(row, columns):
            continue
        line, issue = _line(offset, row, columns)
        if issue is not None:
            issues.append(issue)
        if line is not None:
            lines.append(line)
    return PlanFile(year, tuple(lines), tuple(issues))


def _header_row(grid: list[tuple[object, ...]]) -> int | None:
    for index, row in enumerate(grid):
        if any(_compact(cell) == "кодтп" for cell in row):
            return index
    return None


def _year_of(rows: list[tuple[object, ...]], header: tuple[object, ...]) -> int | None:
    for cell in header:
        found = _year_in(cell)
        if found is not None:
            return found
    for row in rows:
        for cell in row:
            found = _year_in(cell)
            if found is not None:
                return found
    return None


def _year_in(value: object) -> int | None:
    if value is None:
        return None
    found = _YEAR.search(_key(value))
    if found is None:
        return None
    return int(found.group(1))


def _columns(header: tuple[object, ...]) -> dict[str, int]:
    found: dict[str, int] = {}
    names = {
        "наименование покупателя",
        "инн",
        "договор",
        "кодтп",
        "сф тп",
        "адрес тп",
        "тарифная",
        "год",
    }
    for index, cell in enumerate(header):
        key = _key(cell)
        if _compact(cell) == "кодтп":
            found["кодтп"] = index
        elif key == "тарифная" or key.startswith("тарифная "):
            found["тарифная"] = index
        elif _year_volume(key):
            found["год"] = index
        elif key in names or (key in _MONTHS and key not in found):
            found[key] = index
    point = found.get("кодтп")
    if point is not None and point + 1 < len(header) and _key(header[point + 1]) == "":
        found["вид"] = point + 1
    return found


def _line(
    file_row: int,
    row: tuple[object, ...],
    columns: dict[str, int],
) -> tuple[PlanLine | None, PlanIssue | None]:
    try:
        name = _text(row, columns, "наименование покупателя")
        inn = _inn(row, columns)
        contract = _text(row, columns, "договор")
        point = _text(row, columns, "кодтп")
        address = _text(row, columns, "адрес тп")
        region = _REGION_CODES[_key(_cell(row, columns, "сф тп"))]
        stated = _group(_cell(row, columns, "тарифная"))
        volume = volume_of(_cell(row, columns, "год"))
        months = tuple(volume_of(_cell(row, columns, month)) for month in _MONTHS)
    except (KeyError, ValueError, ArithmeticError):
        return None, PlanIssue(file_row, "Строка плана не разбирается.")
    if volume < 0 or any(month < 0 for month in months):
        return None, PlanIssue(file_row, "Отрицательный объём плана.")
    return (
        PlanLine(
            file_row,
            name,
            inn,
            contract,
            point,
            region,
            address,
            _kind(row, columns),
            stated,
            volume,
            months,
        ),
        None,
    )


def _is_total(row: tuple[object, ...], columns: dict[str, int]) -> bool:
    contract = _key(_cell(row, columns, "договор")) if "договор" in columns else ""
    point = _key(_cell(row, columns, "кодтп")) if "кодтп" in columns else ""
    return contract == _TOTAL or point == _TOTAL


def _blank(row: tuple[object, ...], columns: dict[str, int]) -> bool:
    point = _text(row, columns, "кодтп") if "кодтп" in columns else ""
    return point == ""


def _kind(row: tuple[object, ...], columns: dict[str, int]) -> ConsumerKind:
    if "вид" not in columns:
        return ConsumerKind.INDUSTRIAL
    if _key(_cell(row, columns, "вид")) == "комбыт":
        return ConsumerKind.COMMUNAL
    return ConsumerKind.INDUSTRIAL


def _group(value: object) -> Group:
    if isinstance(value, bool) or value is None:
        raise ValueError(value)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int):
        return Group(str(value))
    token = _key(value).replace("a", "а")
    head, _, tail = token.partition(" ")
    if tail.startswith("гр"):
        token = head
    return Group(token)


def _inn(row: tuple[object, ...], columns: dict[str, int]) -> str:
    value = _cell(row, columns, "инн")
    if isinstance(value, bool) or value is None:
        raise ValueError(value)
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(value)
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if text == "":
        raise ValueError(value)
    return text


def _text(row: tuple[object, ...], columns: dict[str, int], key: str) -> str:
    value = _cell(row, columns, key)
    if value is None:
        return ""
    return str(value).strip()


def _cell(row: tuple[object, ...], columns: dict[str, int], key: str) -> object:
    index = columns[key]
    if index >= len(row):
        return None
    return row[index]


def _year_volume(key: str) -> bool:
    if key == "год":
        return True
    year, _, word = key.partition(" ")
    return word == "год" and year.isdigit()


def _compact(value: object) -> str:
    return _key(value).replace(" ", "")


def _key(value: object) -> str:
    if value is None:
        return ""
    folded = str(value).replace("\n", " ").strip().casefold().replace("ё", "е")
    return " ".join(folded.split())
